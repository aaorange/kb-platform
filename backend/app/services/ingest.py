# -*- coding: utf-8 -*-
"""文档解析服务 — 上传格式 → Markdown 文本（供导入引擎消费）

v2 重构后切片/编码/入库由 app/engine 导入图完成，本模块只保留格式解析：
- .md/.txt：编码探测（BOM/UTF-8/UTF-16/gb18030 回退）后原样返回
- .docx：python-docx 按文档顺序解析，标题/列表/表格/图片 → Markdown
- .pdf：由引擎 node_pdf_to_md 处理（MinerU 或 pypdf 本地降级）

docx 图片提取到 work_dir/images/ 并以相对路径引用，
node_md_img 节点会自动做 VL 摘要 + MinIO 托管（与 PDF 图片链路同构）。
"""
import io
import re
from pathlib import Path

SUPPORTED_EXTS = {".md", ".txt", ".docx", ".pdf"}

# 浏览器/VL 模型可直接消费的图片格式（wmf/emf 等 Office 专有格式跳过）
_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
# 兼容英文 "Heading 1" 与中文 Word "标题 1" 两种样式命名
_HEADING_RE = re.compile(r"(?:heading|标题)\s*(\d)")


def parse_file(file_name: str, content: bytes, work_dir: str | None = None) -> str:
    try:
        return _parse_file(file_name, content, work_dir)
    except ValueError:
        raise
    except Exception as e:
        # 损坏的 docx 会抛 python-docx 内部异常，统一转为可读错误
        raise ValueError(f"文档解析失败：{e}")


def _parse_file(file_name: str, content: bytes, work_dir: str | None) -> str:
    ext = "." + file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的格式 {ext}，仅支持 md/txt/docx/pdf")
    if ext in (".md", ".txt"):
        return _decode_text(content)
    if ext == ".docx":
        md = _docx_to_md(content, work_dir)
        if not md.strip():
            raise ValueError("docx 中未提取到任何文本/表格内容（可能是空文档）")
        return md
    raise ValueError("pdf 请交由导入引擎解析（node_pdf_to_md）")


# ===== txt/md：编码探测 =====

def _decode_text(content: bytes) -> str:
    """文本解码：BOM 探测 → utf-8 → gb18030（GBK/GB2312 超集）→ 忽略坏字节兜底

    旧版直接 utf-8 ignore，GBK 中文 txt 会静默丢字导入残缺内容。
    """
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    for enc in ("utf-8", "gb18030"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


# ===== docx：按文档顺序解析（段落/表格/图片交错）=====

def _iter_blocks(doc):
    """按文档顺序产出段落与表格

    doc.paragraphs 与 doc.tables 是两个分离列表，分别遍历会丢失交错顺序
    （表格会全部堆到文末），这里从 body XML 顺序产出。
    """
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _heading_level(para) -> int | None:
    """标题级别（1~6）；兼容中英文样式名，'Title'/'标题' 视为一级"""
    name = (para.style.name or "").strip().lower() if para.style is not None else ""
    m = _HEADING_RE.search(name)
    if m:
        return min(int(m.group(1)), 6)
    if name in ("title", "标题"):
        return 1
    return None


def _num_kind(doc, para) -> tuple[int, int, str] | None:
    """列表项识别：返回 (numId, 缩进级别, bullet|number)，非列表返回 None

    优先读段落 numPr（Word 中英文版本通用，样式名会本地化不可靠）；
    无 numPr 时用样式名兜底（项目符号/编号列表样式）。
    """
    from docx.oxml.ns import qn
    pPr = para._p.pPr
    if pPr is not None and pPr.numPr is not None and pPr.numPr.numId is not None:
        try:
            num_id = int(pPr.numPr.numId.val)
            ilvl = int(pPr.numPr.ilvl.val) if pPr.numPr.ilvl is not None else 0
        except (TypeError, ValueError):
            return None
        return num_id, ilvl, _numbering_fmt(doc, num_id, ilvl)
    name = (para.style.name or "").strip().lower() if para.style is not None else ""
    if "bullet" in name or "项目符号" in name or "列表符号" in name:
        return -1, 0, "bullet"
    if re.search(r"(?:list|列表).*(?:number|编号)", name):
        return -1, 0, "number"
    return None


def _numbering_fmt(doc, num_id: int, ilvl: int) -> str:
    """查 numbering.xml：numId → abstractNumId → ilvl 的 numFmt

    numFmt=bullet 为无序列表，其余（decimal/中文编号等）视为有序列表。
    查不到编号部件时按无序处理（更安全：'-' 不会伪造序号）。
    """
    from docx.oxml.ns import qn
    try:
        numbering = doc.part.numbering_part.element
    except Exception:  # noqa: BLE001 — 文档无编号部件
        return "bullet"
    abs_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(num_id):
            el = num.find(qn("w:abstractNumId"))
            if el is not None:
                abs_id = el.get(qn("w:val"))
            break
    if abs_id is None:
        return "bullet"
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) == abs_id:
            for lvl in abstract.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == str(ilvl):
                    fmt = lvl.find(qn("w:numFmt"))
                    return "number" if fmt is not None and fmt.get(qn("w:val")) != "bullet" else "bullet"
    return "bullet"


def _table_to_md(table) -> str:
    """docx 表格 → Markdown 表格；单元格内换行转 <br>、| 转义；整表为空返回空串"""
    rows, has_content = [], False
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = cell.text.strip().replace("\n", "<br>").replace("|", "\\|")
            if text:
                has_content = True
            cells.append(text or " ")
        rows.append(cells)
    if not rows or not has_content:
        return ""
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def _para_images(para, doc, img_dir: Path | None, seq: int) -> tuple[list[str], int]:
    """提取段落内嵌图片（DrawingML a:blip）到 images/ 目录，返回 md 引用列表"""
    from docx.oxml.ns import qn
    if img_dir is None:
        return [], seq
    refs = []
    for blip in para._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        part = doc.part.related_parts.get(rid)
        blob = getattr(part, "blob", None) if part is not None else None
        if not blob:
            continue
        ext = Path(str(getattr(part, "partname", ""))).suffix.lower()
        if ext not in _IMAGE_EXTS:
            continue
        img_dir.mkdir(parents=True, exist_ok=True)
        fname = f"docx_img_{seq}{ext}"
        (img_dir / fname).write_bytes(blob)
        refs.append(f"![文档图片](images/{fname})")
        seq += 1
    return refs, seq


def _docx_to_md(content: bytes, work_dir: str | None) -> str:
    from docx import Document
    from docx.table import Table

    doc = Document(io.BytesIO(content))
    img_dir = Path(work_dir) / "images" if work_dir else None
    img_seq = 0
    counters: dict[tuple[int, int], int] = {}  # (numId, ilvl) → 有序列表当前序号
    lines: list[str] = []

    for block in _iter_blocks(doc):
        if isinstance(block, Table):
            md = _table_to_md(block)
            if md:
                lines.extend(["", md, ""])
            continue

        text = block.text.strip()
        img_refs, img_seq = _para_images(block, doc, img_dir, img_seq)

        if not text:
            if img_refs:  # 纯图片段落
                lines.extend(["", *img_refs, ""])
            continue

        level = _heading_level(block)
        if level:
            lines.extend(["", "#" * level + " " + text, ""])
            continue

        info = _num_kind(doc, block)
        if info is not None:
            num_id, ilvl, kind = info
            if kind == "number":
                counters[(num_id, ilvl)] = counters.get((num_id, ilvl), 0) + 1
                marker = f"{counters[(num_id, ilvl)]}. "
            else:
                marker = "- "
            lines.append("  " * ilvl + marker + text)
            continue

        if img_refs:  # 图文混排段落：文本在前图片在后
            lines.extend(["", text, *img_refs])
        else:
            lines.append(text)

    return "\n".join(lines).strip()
